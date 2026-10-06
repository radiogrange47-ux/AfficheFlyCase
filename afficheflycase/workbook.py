"""Import and validate the AfficheFlyCase Excel workbook format."""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet


SHEET_HEADERS = {
    "Listing Materiel": {
        2: "ELEMENT",
        3: "POSITION",
        4: "QTE",
        5: "SPARE",
        6: "FLYCASE",
    },
    "Listing FlyCase": {
        2: "ID",
        3: "Type2",
        4: "Couleur",
        5: "Largeur",
        6: "Longeur",
        7: "Hauteur",
        8: "Empatement",
        9: "Cubage",
        10: "Tip",
        11: "Gerbable",
        12: "Commentaire",
    },
}

FLYCASE_ID_PATTERN = re.compile(r"(?<![A-Z0-9])([A-Z]{1,8}\d{1,8})(?![A-Z0-9])", re.IGNORECASE)


class WorkbookFormatError(ValueError):
    """Raised when a workbook does not match the documented template."""


@dataclass(frozen=True)
class MaterialItem:
    element: str
    position: str
    quantity: str
    spare: str
    source_row: int


@dataclass
class FlyCase:
    identifier: str
    case_type: str
    color: str
    width: str
    length: str
    height: str
    footprint: str
    volume: str
    tip: str
    stackable: str
    comment: str
    source_row: int
    materials: list[MaterialItem] = field(default_factory=list)

    @property
    def dimensions(self) -> str:
        dimensions = [value for value in (self.width, self.length, self.height) if value]
        return " × ".join(dimensions) if dimensions else "À vérifier"

    @property
    def tip_label(self) -> str:
        if self.tip == "NON":
            return "NE PAS BASCULER"
        if self.tip == "OK":
            return "BASCULABLE"
        return "STATUT À VÉRIFIER"

    @property
    def stackable_label(self) -> str:
        if self.stackable == "OK":
            return "GERBABLE"
        if self.stackable == "NON":
            return "NON GERBABLE"
        return "STATUT À VÉRIFIER"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def _normalized_id(value: str) -> str:
    return value.strip().upper()


def _merged_value(sheet: Worksheet, row: int, column: int) -> Any:
    cell = sheet.cell(row=row, column=column)
    if not isinstance(cell, MergedCell):
        return cell.value
    for merged_range in sheet.merged_cells.ranges:
        if cell.coordinate in merged_range:
            return sheet.cell(merged_range.min_row, merged_range.min_col).value
    return None


def _cached_value(
    formulas_sheet: Worksheet, values_sheet: Worksheet, row: int, column: int, sheet_name: str
) -> Any:
    formula = _merged_value(formulas_sheet, row, column)
    value = _merged_value(values_sheet, row, column)
    if isinstance(formula, str) and formula.startswith("=") and value is None:
        address = formulas_sheet.cell(row=row, column=column).coordinate
        raise WorkbookFormatError(
            f"Formule sans résultat enregistré dans « {sheet_name} »!{address}. "
            "Ouvrez le classeur dans Excel, recalculez-le, enregistrez-le puis réessayez."
        )
    return value


def _validate_sheet_names(workbook) -> None:
    expected = list(SHEET_HEADERS)
    actual = workbook.sheetnames
    for index, expected_name in enumerate(expected):
        if index >= len(actual):
            raise WorkbookFormatError(
                f"Onglet « {expected_name} » manquant. Le classeur doit commencer par "
                "« Listing Materiel », puis « Listing FlyCase »."
            )
        if actual[index].strip().casefold() != expected_name.casefold():
            raise WorkbookFormatError(
                f"Le premier onglet attendu à la position {index + 1} est "
                f"« {expected_name} », mais le fichier contient « {actual[index]} »."
            )


def _validate_headers(sheet: Worksheet, expected: dict[int, str]) -> None:
    for column, expected_header in expected.items():
        actual = _text(sheet.cell(row=2, column=column).value)
        if actual.casefold() != expected_header.casefold():
            address = sheet.cell(row=2, column=column).coordinate
            displayed = actual if actual else "cellule vide"
            raise WorkbookFormatError(
                f"En-tête incorrect dans « {sheet.title} »!{address} : "
                f"attendu « {expected_header} », trouvé « {displayed} »."
            )


def _extract_case_ids(description: str) -> list[str]:
    return [_normalized_id(match) for match in FLYCASE_ID_PATTERN.findall(description.upper())]


def _expand_material_rows(sheet: Worksheet, row: int) -> tuple[str, str, str, str, str]:
    values = [_text(_merged_value(sheet, row, column)) for column in range(2, 7)]
    element, position, quantity, spare, flycase = values
    return element, position, quantity, spare, flycase


def load_flycases(path: str | Path) -> list[FlyCase]:
    """Validate and load all fly cases from the first two workbook sheets."""
    workbook_path = Path(path)
    if workbook_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise WorkbookFormatError("Format non pris en charge. Sélectionnez un fichier .xlsx ou .xlsm.")
    if not workbook_path.is_file():
        raise WorkbookFormatError("Le fichier Excel sélectionné est introuvable.")

    try:
        formulas_book = load_workbook(workbook_path, read_only=False, data_only=False, keep_vba=False)
        values_book = load_workbook(workbook_path, read_only=False, data_only=True, keep_vba=False)
    except PermissionError as error:
        raise WorkbookFormatError(
            "Impossible d’accéder au fichier Excel. Fermez-le dans Excel ou vérifiez les droits "
            "d’accès, puis réessayez."
        ) from error
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        raise WorkbookFormatError(
            f"Impossible de lire le fichier Excel : {error}. Vérifiez qu’il n’est pas corrompu "
            "et qu’il s’agit bien d’un classeur Excel pris en charge."
        ) from error

    try:
        _validate_sheet_names(formulas_book)
        material_formula_sheet = formulas_book["Listing Materiel"]
        flycase_formula_sheet = formulas_book["Listing FlyCase"]
        material_values_sheet = values_book["Listing Materiel"]
        flycase_values_sheet = values_book["Listing FlyCase"]
        _validate_headers(material_formula_sheet, SHEET_HEADERS["Listing Materiel"])
        _validate_headers(flycase_formula_sheet, SHEET_HEADERS["Listing FlyCase"])

        cases: dict[str, FlyCase] = {}
        for row in range(3, flycase_formula_sheet.max_row + 1):
            values = [
                _cached_value(flycase_formula_sheet, flycase_values_sheet, row, column, "Listing FlyCase")
                for column in range(2, 13)
            ]
            if not any(_text(value) for value in values):
                continue
            identifier = _normalized_id(_text(values[0]))
            if not identifier:
                raise WorkbookFormatError(
                    f"ID manquant dans « Listing FlyCase »!B{row}. Corrigez la ligne ou supprimez-la."
                )
            if identifier in cases:
                raise WorkbookFormatError(
                    f"ID en double « {identifier} » dans « Listing FlyCase »!B{row}."
                )
            statuses = {}
            for column, label, raw in (
                ("J", "Tip", _text(values[8])),
                ("K", "Gerbable", _text(values[9])),
            ):
                status = raw.upper()
                if status not in {"", "OK", "NON"}:
                    raise WorkbookFormatError(
                        f"Valeur « {raw} » non reconnue dans « Listing FlyCase »!{column}{row} "
                        f"({label}). Valeurs acceptées : OK, NON ou cellule vide."
                    )
                statuses[label] = status
            cases[identifier] = FlyCase(
                identifier=identifier,
                case_type=_text(values[1]),
                color=_text(values[2]),
                width=_text(values[3]),
                length=_text(values[4]),
                height=_text(values[5]),
                footprint=_text(values[6]),
                volume=_text(values[7]),
                tip=statuses["Tip"],
                stackable=statuses["Gerbable"],
                comment=_text(values[10]),
                source_row=row,
            )

        if not cases:
            raise WorkbookFormatError(
                "Aucune fiche avec un ID n’a été trouvée dans « Listing FlyCase »."
            )

        for row in range(3, material_formula_sheet.max_row + 1):
            raw = [
                _cached_value(material_formula_sheet, material_values_sheet, row, column, "Listing Materiel")
                for column in range(2, 7)
            ]
            if not any(_text(value) for value in raw):
                continue
            element, position, quantity, spare, description = _expand_material_rows(
                material_values_sheet, row
            )
            # _expand_material_rows reads cached cells and applies merged values.
            if not _text(raw[4]):
                if any(_text(value) for value in raw[:4]):
                    raise WorkbookFormatError(
                        f"Fly case manquant dans « Listing Materiel »!F{row}."
                    )
                continue
            identifiers = _extract_case_ids(_text(raw[4]))
            if not identifiers:
                raise WorkbookFormatError(
                    f"Aucun ID de fly case reconnu dans « Listing Materiel »!F{row} "
                    f"(« {_text(raw[4])} »)."
                )
            material = MaterialItem(
                element=element or _text(raw[0]),
                position=position or _text(raw[1]),
                quantity=quantity or _text(raw[2]),
                spare=spare or _text(raw[3]),
                source_row=row,
            )
            for identifier in identifiers:
                if identifier not in cases:
                    raise WorkbookFormatError(
                        f"ID « {identifier} » de « Listing Materiel »!F{row} "
                        "absent de l’onglet « Listing FlyCase »."
                    )
                cases[identifier].materials.append(material)
        return list(cases.values())
    finally:
        formulas_book.close()
        values_book.close()

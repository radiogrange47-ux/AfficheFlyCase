import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from afficheflycase.workbook import WorkbookFormatError, load_flycases


def make_workbook(path: Path, *, sheet_names=None, bad_header=False, bad_status=False) -> None:
    workbook = Workbook()
    material = workbook.active
    material.title = (sheet_names or ["Listing Materiel", "Listing FlyCase"])[0]
    material.append(["", "", "", "", "", ""])
    material.append(["", "WRONG" if bad_header else "ELEMENT", "POSITION", "QTE", "SPARE", "FLYCASE"])
    material.merge_cells("B3:B4")
    material.merge_cells("C3:C4")
    material.merge_cells("D3:D4")
    material.merge_cells("E3:E4")
    material["B3"] = "Lampe"
    material["C3"] = "Scène"
    material["D3"] = 2
    material["E3"] = "-"
    material["F3"] = "LX01 - Fly Case - Noir"
    material["F4"] = "LX02 - Fly Case - Bleu"

    flycases = workbook.create_sheet((sheet_names or ["Listing Materiel", "Listing FlyCase"])[1])
    headers = ["", "ID", "Type2", "Couleur", "Largeur", "Longeur", "Hauteur",
               "Empatement", "Cubage", "Tip", "Gerbable", "Commentaire"]
    flycases.append([""] * len(headers))
    flycases.append(["", "ID", "Type2", "Couleur", "Largeur", "Longeur", "Hauteur",
                     "Empatement", "Cubage", "Tip", "Gerbable", "Commentaire"])
    flycases.append(["", "LX01", "Fly Case", "Noir", "120 cm", "60 cm", "63 cm",
                     "0,7 m²", "0,45 m³", "OK", "OK", "Fragile"])
    flycases.append(["", "LX02", "Fly Case", "Bleu", "50 cm", "40 cm", "30 cm",
                     "-", "0,06 m³", "NON" if not bad_status else "MAYBE", "NON", ""])
    workbook.save(path)


class WorkbookImportTests(unittest.TestCase):
    def test_loads_cases_and_expands_merged_material_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "listing.xlsx"
            make_workbook(path)
            cases = load_flycases(path)
        self.assertEqual([case.identifier for case in cases], ["LX01", "LX02"])
        self.assertEqual(len(cases[0].materials), 1)
        self.assertEqual(len(cases[1].materials), 1)
        self.assertEqual(cases[1].materials[0].element, "Lampe")
        self.assertEqual(cases[0].dimensions, "120 cm × 60 cm × 63 cm")
        self.assertEqual(cases[1].tip_label, "NE PAS BASCULER")

    def test_rejects_incorrect_header_with_cell_location(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wrong.xlsx"
            make_workbook(path, bad_header=True)
            with self.assertRaisesRegex(WorkbookFormatError, r"Listing Materiel.*B2"):
                load_flycases(path)

    def test_rejects_wrong_sheet_order(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wrong.xlsx"
            make_workbook(path, sheet_names=["Other", "Listing FlyCase"])
            with self.assertRaisesRegex(WorkbookFormatError, "position 1"):
                load_flycases(path)

    def test_rejects_unknown_handling_status(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wrong.xlsx"
            make_workbook(path, bad_status=True)
            with self.assertRaisesRegex(WorkbookFormatError, r"J4.*Tip"):
                load_flycases(path)

    def test_rejects_unmatched_material_case_id(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wrong.xlsx"
            make_workbook(path)
            from openpyxl import load_workbook

            workbook = load_workbook(path)
            workbook["Listing Materiel"]["F3"] = "LX99 - Fly Case"
            workbook.save(path)
            with self.assertRaisesRegex(WorkbookFormatError, r"LX99.*F3"):
                load_flycases(path)


if __name__ == "__main__":
    unittest.main()
